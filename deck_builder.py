"""
deck_builder.py - FINCH Presentation Generator
Generates clean, executive-styled PowerPoint (.pptx) decks with custom color themes,
metric scorecards, visual bar charts, and structured action cards.
"""

import io
import pandas as pd

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_executive_deck(facts, currency="₦", filename="Dataset", scenarios=None):
    """
    Generates an 8-slide executive presentation byte stream using native PPTX elements
    matching FINCH's dark/light executive theme.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # FINCH Theme Colors
    theme_colors = {
        "bg_dark": RGBColor(9, 13, 22),       # Deep Navy #090D16
        "bg_light": RGBColor(255, 255, 255),   # Clean White
        "card_bg": RGBColor(248, 250, 252),    # Soft Off-White
        "primary": RGBColor(15, 23, 42),       # Dark Slate #0F172A
        "accent": RGBColor(15, 157, 143),      # Teal Accent #0F9D8F
        "success": RGBColor(16, 185, 129),     # Emerald Green #10B981
        "gold": RGBColor(244, 184, 96),        # Gold #F4B860
        "text_muted": RGBColor(100, 116, 139), # Muted Gray
    }

    totals = facts["totals"]
    meta = facts.get("metadata", {})
    rbc = facts.get("revenue_by_category", {})

    def set_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, category: str, title: str, subtitle: str = ""):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = theme_colors["accent"]

        txBox_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf_title = txBox_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = theme_colors["primary"]

        if subtitle:
            txBox_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.4))
            tf_sub = txBox_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 1: Title Slide
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide1, theme_colors["bg_dark"])

    txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "FINCH EXECUTIVE REPORT"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = theme_colors["gold"]
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = f"Business Intelligence Audit: {filename}"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(248, 250, 252)
    p1.space_after = Pt(16)

    margin_val = totals.get('profit_margin')
    margin_str = f"{margin_val:.1f}%" if margin_val is not None else "N/A"
    p2 = tf.add_paragraph()
    p2.text = f"Automated financial health audit, growth scenarios, and executive recommendations. • Net Margin: {margin_str}"
    p2.font.size = Pt(16)
    p2.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 2: Financial Summary
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide2, theme_colors["bg_light"])
    add_header(slide2, "Core Metrics", "Executive Financial Summary", "High-level summary of total top-line revenue, operating costs, and net profit.")

    cost_known = totals.get("cost_known", False)
    rev_str = f"{currency}{totals.get('total_revenue', 0):,.2f}"
    cost_str = f"{currency}{totals.get('total_costs', 0):,.2f}" if cost_known else "N/A"
    profit_str = f"{currency}{totals.get('net_profit', 0):,.2f}" if cost_known else "N/A"

    metrics = [
        {"label": "TOTAL REVENUE", "value": rev_str, "sub": "Gross Inflow", "color": theme_colors["primary"], "bg": theme_colors["card_bg"]},
        {"label": "TOTAL COSTS", "value": cost_str, "sub": "Operating Expenses", "color": theme_colors["primary"], "bg": theme_colors["card_bg"]},
        {"label": "NET PROFIT", "value": profit_str, "sub": f"{margin_str} Net Margin", "color": theme_colors["success"], "bg": theme_colors["primary"], "dark": True}
    ]

    card_width = Inches(3.6)
    card_height = Inches(3.5)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    start_y = Inches(2.2)

    for i, m in enumerate(metrics):
        x = start_x + i * (card_width + gap)
        shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, start_y, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = m["bg"]
        shape.line.color.rgb = theme_colors["card_bg"] if m.get("dark") else RGBColor(226, 232, 240)

        tf_card = shape.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.3)
        tf_card.margin_top = Inches(0.4)

        p_lbl = tf_card.paragraphs[0]
        p_lbl.text = m["label"]
        p_lbl.font.size = Pt(12)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = theme_colors["text_muted"] if not m.get("dark") else RGBColor(148, 163, 184)
        p_lbl.space_after = Pt(16)

        p_val = tf_card.add_paragraph()
        p_val.text = m["value"]
        p_val.font.size = Pt(28)
        p_val.font.bold = True
        p_val.font.color.rgb = m["color"]
        p_val.space_after = Pt(12)

        p_sub = tf_card.add_paragraph()
        p_sub.text = m["sub"]
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = theme_colors["accent"] if not m.get("dark") else theme_colors["success"]

    # ========================================== SLIDE 3: Financial Overview & Core Metrics
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide3, theme_colors["bg_light"])
    add_header(slide3, "Performance Breakdown", "Financial Overview & Core Metrics", "Detailed breakdown of profitability and operational efficiency.")

    box3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
    box3.fill.solid()
    box3.fill.fore_color.rgb = theme_colors["card_bg"]
    box3.line.color.rgb = RGBColor(226, 232, 240)

    tf3 = box3.text_frame
    tf3.margin_left = Inches(0.5)
    tf3.margin_top = Inches(0.5)
    
    p3_0 = tf3.paragraphs[0]
    p3_0.text = "EXECUTIVE LEDGER METRICS"
    p3_0.font.size = Pt(14)
    p3_0.font.bold = True
    p3_0.font.color.rgb = theme_colors["accent"]
    p3_0.space_after = Pt(16)

    p3_1 = tf3.add_paragraph()
    p3_1.text = f"• Gross Top-Line Revenue: {rev_str}\n" \
               f"• Total Recorded Costs: {cost_str}\n" \
               f"• Bottom-Line Net Profit: {profit_str}\n" \
               f"• Operating Margin Efficiency: {margin_str}"
    p3_1.font.size = Pt(18)
    p3_1.font.color.rgb = theme_colors["primary"]

    # ========================================== SLIDE 4: Top Categories (Visual Bar Chart)
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide4, theme_colors["bg_light"])
    add_header(slide4, "Category Breakdown", "Top Revenue Categories & Concentration", "Comparative distribution of gross sales performance across categories.")

    if rbc:
        max_rev = max(list(rbc.values())) if rbc else 1
        bar_start_y = Inches(2.2)
        row_height = Inches(1.1)

        for i, (cat_name, rev) in enumerate(list(rbc.items())[:4]):
            y = bar_start_y + i * row_height
            pct = (rev / totals.get("total_revenue", 1)) * 100 if totals.get("total_revenue", 0) > 0 else 0

            txBox = slide4.shapes.add_textbox(Inches(0.8), y, Inches(11.5), Inches(0.35))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = f"{cat_name}   —   {currency}{rev:,.2f} ({pct:.1f}%)"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = theme_colors["primary"]

            track = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y + Inches(0.4), Inches(11.5), Inches(0.3))
            track.fill.solid()
            track.fill.fore_color.rgb = RGBColor(226, 232, 240)
            track.line.fill.background()

            fill_width = Inches(11.5 * (rev / max_rev)) if max_rev > 0 else Inches(0.5)
            fill_bar = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y + Inches(0.4), fill_width, Inches(0.3))
            fill_bar.fill.solid()
            fill_bar.fill.fore_color.rgb = theme_colors["accent"]
            fill_bar.line.fill.background()
    else:
        txBox = slide4.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.0))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "No category mapping selected for breakdown."
        p.font.size = Pt(16)
        p.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 5: Data Health & Quality Check
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide5, theme_colors["bg_light"])
    add_header(slide5, "Data Governance", "Data Health & Outlier Quality Check", "Automated screening for null cells, duplicate records, and statistical anomalies.")

    circle = slide5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(2.5), Inches(3.2), Inches(3.2))
    circle.fill.solid()
    circle.fill.fore_color.rgb = theme_colors["card_bg"]
    circle.line.color.rgb = theme_colors["success"]
    circle.line.width = Pt(3)

    tf_c = circle.text_frame
    tf_c.margin_top = Inches(0.9)
    p_c1 = tf_c.paragraphs[0]
    health_score = meta.get('data_health_index', 100)
    p_c1.text = f"{health_score}%"
    p_c1.font.size = Pt(44)
    p_c1.font.bold = True
    p_c1.font.color.rgb = theme_colors["success"]
    p_c1.alignment = PP_ALIGN.CENTER

    p_c2 = tf_c.add_paragraph()
    p_c2.text = "Health Score"
    p_c2.font.size = Pt(13)
    p_c2.font.color.rgb = theme_colors["text_muted"]
    p_c2.alignment = PP_ALIGN.CENTER

    anomalies_list = facts.get("anomalies", [])
    rows = [
        {"title": "Duplicate Records", "desc": f"{meta.get('duplicate_rows', 0):,} duplicate rows identified in primary audit", "val": str(meta.get('duplicate_rows', 0))},
        {"title": "Missing Data Cells", "desc": f"{meta.get('missing_cells', 0):,} missing cells found across active columns", "val": str(meta.get('missing_cells', 0))},
        {"title": "Statistical Anomaly Detection", "desc": f"Detected {len(anomalies_list)} transaction spike(s) for verification", "val": str(len(anomalies_list))}
    ]

    for i, r in enumerate(rows):
        y = Inches(2.4) + i * Inches(1.2)
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.0), y, Inches(7.3), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = theme_colors["card_bg"]
        box.line.color.rgb = RGBColor(226, 232, 240)

        tf_r = box.text_frame
        tf_r.margin_left = Inches(0.3)
        tf_r.margin_top = Inches(0.2)
        p_t = tf_r.paragraphs[0]
        p_t.text = r["title"]
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = theme_colors["primary"]

        p_d = tf_r.add_paragraph()
        p_d.text = r["desc"]
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 6: What-If Growth Scenarios
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide6, theme_colors["bg_light"])
    add_header(slide6, "Sensitivity Modelling", "What-If Growth & Sensitivity Scenarios", "Simulated business performance under varying price and volume adjustments.")

    if scenarios and len(scenarios) > 0:
        card_w = Inches(3.6)
        card_h = Inches(4.5)
        gap_s = Inches(0.4)
        start_s_x = Inches(0.8)

        for i, sc in enumerate(scenarios[:3]):
            x = start_s_x + i * (card_w + gap_s)
            box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), card_w, card_h)
            box.fill.solid()
            box.fill.fore_color.rgb = theme_colors["card_bg"]
            box.line.color.rgb = theme_colors["accent"]

            tf_s = box.text_frame
            tf_s.margin_left = Inches(0.3)
            tf_s.margin_top = Inches(0.3)
            tf_s.word_wrap = True

            p_sn = tf_s.paragraphs[0]
            p_sn.text = sc.get("scenario_name", f"Scenario {i+1}").upper()
            p_sn.font.size = Pt(14)
            p_sn.font.bold = True
            p_sn.font.color.rgb = theme_colors["accent"]
            p_sn.space_after = Pt(16)

            sc_rev = sc.get('revenue', 0)
            sc_cost = sc.get('costs', 0)
            sc_prof = sc.get('net_profit', 0)
            sc_marg = sc.get('margin', 0)

            p_st = tf_s.add_paragraph()
            p_st.text = f"• Simulated Rev: {currency}{sc_rev:,.2f}\n" \
                        f"• Simulated Costs: {currency}{sc_cost:,.2f}\n" \
                        f"• Net Profit: {currency}{sc_prof:,.2f}\n" \
                        f"• Profit Margin: {sc_marg:.1f}%"
            p_st.font.size = Pt(14)
            p_st.font.color.rgb = theme_colors["primary"]
    else:
        box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
        box.fill.solid()
        box.fill.fore_color.rgb = theme_colors["card_bg"]
        box.line.color.rgb = RGBColor(226, 232, 240)
        tf_nos = box.text_frame
        tf_nos.margin_left = Inches(0.5)
        tf_nos.margin_top = Inches(0.5)
        p = tf_nos.paragraphs[0]
        p.text = "No saved scenario models present for comparison. Configure price & volume adjustments in the What-If tab."
        p.font.size = Pt(16)
        p.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 7: Strategic Action Plan
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide7, theme_colors["bg_light"])
    add_header(slide7, "Execution Roadmap", "FINCH Strategic Action Plan", "Prioritized recommendations to sustain growth and operational efficiency.")

    actions = [
        {"num": "01", "title": "Margin Optimization", "desc": "Audit high-cost lines and optimize variable per-unit expenses."},
        {"num": "02", "title": "Category Scaling", "desc": "Scale channel investment into top-performing regional drivers."},
        {"num": "03", "title": "Anomaly Resolution", "desc": "Verify outlier transactions to preserve ledger consistency."},
        {"num": "04", "title": "Scenario Targets", "desc": "Deploy sensitivity forecasts to protect bottom-line margins."}
    ]

    card_w = Inches(2.7)
    card_h = Inches(4.2)
    gap_a = Inches(0.3)
    start_a_x = Inches(0.8)

    for i, act in enumerate(actions):
        x = start_a_x + i * (card_w + gap_a)
        box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = theme_colors["card_bg"]
        box.line.color.rgb = RGBColor(226, 232, 240)

        tf_a = box.text_frame
        tf_a.margin_left = Inches(0.25)
        tf_a.margin_top = Inches(0.3)
        tf_a.word_wrap = True

        p_num = tf_a.paragraphs[0]
        p_num.text = act["num"]
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = theme_colors["accent"]
        p_num.space_after = Pt(14)

        p_at = tf_a.add_paragraph()
        p_at.text = act["title"]
        p_at.font.size = Pt(15)
        p_at.font.bold = True
        p_at.font.color.rgb = theme_colors["primary"]
        p_at.space_after = Pt(10)

        p_ad = tf_a.add_paragraph()
        p_ad.text = act["desc"]
        p_ad.font.size = Pt(12)
        p_ad.font.color.rgb = theme_colors["text_muted"]

    # ========================================== SLIDE 8: Closing Slide
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide8, theme_colors["bg_dark"])

    txBox8 = slide8.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.3), Inches(3.0))
    tf8 = txBox8.text_frame
    tf8.word_wrap = True

    p8_1 = tf8.paragraphs[0]
    p8_1.text = "Audit Complete & Verified"
    p8_1.font.size = Pt(38)
    p8_1.font.bold = True
    p8_1.font.color.rgb = RGBColor(248, 250, 252)
    p8_1.alignment = PP_ALIGN.CENTER
    p8_1.space_after = Pt(16)

    p8_2 = tf8.add_paragraph()
    p8_2.text = "FINCH Business Intelligence Workspace"
    p8_2.font.size = Pt(18)
    p8_2.font.color.rgb = theme_colors["success"]
    p8_2.alignment = PP_ALIGN.CENTER

    # Output to Byte Buffer
    bio = io.BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio.getvalue()